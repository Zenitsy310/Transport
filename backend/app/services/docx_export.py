from __future__ import annotations

import io
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def build_docx(answer_text: str, request_id: str | None = None) -> bytes:
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)

    header = document.add_paragraph("ГУП «Мосгортранс»")
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].bold = True

    meta = document.add_paragraph(f"Дата: {date.today().isoformat()}")
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if request_id:
        rid = document.add_paragraph(f"Идентификатор черновика: {request_id}")
        rid.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rid.runs[0].font.size = Pt(10)
        rid.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    document.add_paragraph()
    for block in answer_text.split("\n"):
        paragraph = document.add_paragraph(block if block else "")
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
