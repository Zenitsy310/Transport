from __future__ import annotations

import asyncio
import io
import logging
import zipfile
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET

from app.config import Settings

logger = logging.getLogger(__name__)

IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png"}
MIN_DIGITAL_CHARS = 40
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class OCRResult:
    text: str
    method: str
    success: bool
    error: str | None = None


class OCRAdapter(ABC):
    @abstractmethod
    async def extract_text(self, payload: bytes, filename: str) -> OCRResult:
        raise NotImplementedError


class StubOCRAdapter(OCRAdapter):
    """Deterministic fallback for environments without PaddleOCR."""

    async def extract_text(self, payload: bytes, filename: str) -> OCRResult:
        suffix = Path(filename).suffix.lower()
        if suffix == ".txt":
            return OCRResult(payload.decode("utf-8", errors="replace"), "txt", True)
        if suffix in {".docx", ".doc"}:
            return await asyncio.to_thread(_extract_docx, payload)
        if suffix == ".pdf":
            digital = await asyncio.to_thread(_extract_pdf_digital, payload)
            if digital.success and len(digital.text.strip()) >= MIN_DIGITAL_CHARS:
                return digital
            return OCRResult(
                digital.text or "[stub-ocr] Изображение/скан PDF: установите PaddleOCR для распознавания.",
                "stub-pdf",
                True,
            )
        if suffix in IMAGE_SUFFIXES:
            return OCRResult(
                "[stub-ocr] Скан обращения. Для распознавания изображений включите OCR_PROVIDER=paddle.",
                "stub-image",
                True,
            )
        return OCRResult("", "unsupported", False, f"Неподдерживаемый формат: {suffix}")


class PaddleOCRAdapter(OCRAdapter):
    def __init__(self) -> None:
        self._engine = None

    def _get_engine(self):
        if self._engine is None:
            from paddleocr import PaddleOCR

            self._engine = PaddleOCR(use_angle_cls=True, lang="ru", show_log=False)
        return self._engine

    async def extract_text(self, payload: bytes, filename: str) -> OCRResult:
        suffix = Path(filename).suffix.lower()
        if suffix == ".txt":
            return OCRResult(payload.decode("utf-8", errors="replace"), "txt", True)
        if suffix in {".docx", ".doc"}:
            return await asyncio.to_thread(_extract_docx, payload)
        if suffix == ".pdf":
            digital = await asyncio.to_thread(_extract_pdf_digital, payload)
            if digital.success and len(digital.text.strip()) >= MIN_DIGITAL_CHARS:
                return digital
            return await asyncio.to_thread(self._ocr_pdf, payload)
        if suffix in IMAGE_SUFFIXES:
            return await asyncio.to_thread(self._ocr_image_bytes, payload)
        return OCRResult("", "unsupported", False, f"Неподдерживаемый формат: {suffix}")

    def _ocr_image_bytes(self, payload: bytes) -> OCRResult:
        from PIL import Image

        image = Image.open(io.BytesIO(payload)).convert("RGB")
        return self._run_ocr(image, "paddle-image")

    def _ocr_pdf(self, payload: bytes) -> OCRResult:
        import fitz

        doc = fitz.open(stream=payload, filetype="pdf")
        pages: list[str] = []
        try:
            for page in doc:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2))
                from PIL import Image

                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                result = self._run_ocr(image, "paddle-pdf")
                if result.text:
                    pages.append(result.text)
        finally:
            doc.close()
        text = "\n\n".join(pages).strip()
        if not text:
            return OCRResult("", "paddle-pdf", False, "PaddleOCR не извлёк текст из PDF")
        return OCRResult(text, "paddle-pdf", True)

    def _run_ocr(self, image, method: str) -> OCRResult:
        import numpy as np

        engine = self._get_engine()
        result = engine.ocr(np.array(image), cls=True)
        lines: list[str] = []
        if result:
            for page in result:
                if not page:
                    continue
                for item in page:
                    if item and len(item) >= 2 and item[1]:
                        lines.append(str(item[1][0]))
        text = "\n".join(lines).strip()
        if not text:
            return OCRResult("", method, False, "PaddleOCR не распознал текст")
        return OCRResult(text, method, True)


def _extract_docx(payload: bytes) -> OCRResult:
    if payload.startswith(OLE_MAGIC):
        return OCRResult(
            "",
            "doc",
            False,
            "Формат Word 97–2003 (.doc) не поддерживается. Сохраните обращение как .docx и загрузите снова.",
        )
    try:
        result = _extract_docx_library(payload)
        if result.success and result.text.strip():
            return result
    except Exception as exc:
        logger.warning("docx_library_failed", extra={"event": "docx_library_failed", "error": str(exc)})
    try:
        text = _extract_docx_xml(payload)
        if text:
            return OCRResult(text, "docx-xml", True)
    except Exception as exc:
        logger.warning("docx_xml_failed", extra={"event": "docx_xml_failed", "error": str(exc)})
        return OCRResult("", "docx", False, f"Не удалось прочитать Word-файл: {exc}")
    return OCRResult("", "docx", False, "Word-файл не содержит извлекаемого текста")


def _extract_docx_library(payload: bytes) -> OCRResult:
    from docx import Document

    document = Document(io.BytesIO(payload))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            parts.append(paragraph.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                if paragraph.text.strip():
                    parts.append(paragraph.text.strip())
    text = "\n".join(parts).strip()
    if not text:
        return OCRResult("", "docx", False, "DOCX не содержит текста")
    return OCRResult(text, "docx", True)


def _extract_docx_xml(payload: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = [
            name
            for name in archive.namelist()
            if name.startswith("word/") and name.endswith(".xml") and not name.endswith(".rels")
        ]
        chunks: list[str] = []
        for name in names:
            if not name.startswith(("word/document", "word/header", "word/footer", "word/footnotes", "word/endnotes")):
                continue
            root = ET.fromstring(archive.read(name))
            paragraphs: list[str] = []
            for paragraph in root.iter(f"{W_NS}p"):
                line = "".join(node.text or "" for node in paragraph.iter(f"{W_NS}t")).strip()
                if line:
                    paragraphs.append(line)
            if paragraphs:
                chunks.append("\n".join(paragraphs))
    return "\n".join(chunk.strip() for chunk in chunks if chunk.strip()).strip()


def _extract_pdf_digital(payload: bytes) -> OCRResult:
    import fitz

    doc = fitz.open(stream=payload, filetype="pdf")
    try:
        text = "\n".join(page.get_text() for page in doc).strip()
    finally:
        doc.close()
    if not text:
        return OCRResult("", "pdf-digital", False, "В PDF нет цифрового текстового слоя")
    return OCRResult(text, "pdf-digital", True)


def build_ocr_adapter(settings: Settings) -> OCRAdapter:
    if settings.ocr_provider == "stub":
        logger.info("ocr_adapter_selected", extra={"event": "ocr_adapter_selected", "provider": "stub"})
        return StubOCRAdapter()
    try:
        import paddleocr  # noqa: F401

        logger.info("ocr_adapter_selected", extra={"event": "ocr_adapter_selected", "provider": "paddle"})
        return PaddleOCRAdapter()
    except Exception as exc:  # pragma: no cover - environment dependent
        logger.warning(
            "ocr_fallback_stub",
            extra={"event": "ocr_fallback_stub", "provider": "stub", "error": str(exc)},
        )
        return StubOCRAdapter()
